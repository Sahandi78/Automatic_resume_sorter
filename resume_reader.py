
from pathlib import Path

import pymupdf
import pymupdf4llm
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


# --------------------------------------------------
# Detect whether a PDF page has two text columns
# --------------------------------------------------

def detect_two_columns(page):
    """
    Estimate whether a page contains two distinct columns.

    Returns:
        (True, boundary) if two columns are detected.
        (False, None) otherwise.

    This is a heuristic and may not detect every layout.
    """

    page_width = page.rect.width
    blocks = page.get_text("blocks")

    # Ignore empty blocks and very wide blocks, such as
    # headings that span most of the page.
    usable_blocks = []

    for block in blocks:
        x0, y0, x1, y1, text = block[:5]

        if not text.strip():
            continue

        width = x1 - x0

        if width < page_width * 0.75:
            center_x = (x0 + x1) / 2
            usable_blocks.append(center_x)

    # Require enough blocks to make a useful decision.
    if len(usable_blocks) < 6:
        return False, None

    centers = sorted(usable_blocks)

    # Find the largest horizontal gap between block centers.
    gaps = [
        (centers[i + 1] - centers[i], i)
        for i in range(len(centers) - 1)
    ]

    largest_gap, gap_index = max(gaps, key=lambda item: item[0])

    left_count = gap_index + 1
    right_count = len(centers) - left_count

    # Reject small gaps and splits with too few blocks
    # on either side.
    if largest_gap < page_width * 0.15:
        return False, None

    if left_count < 3 or right_count < 3:
        return False, None

    boundary = (
        centers[gap_index] + centers[gap_index + 1]
    ) / 2

    return True, boundary


# --------------------------------------------------
# Extract a PDF using coordinate-based columns
# --------------------------------------------------

def extract_pdf_columns(file_path: Path) -> str:
    """
    Extract PDF text while grouping blocks by column.

    Full-width blocks are preserved separately so the
    LLM can use them to reconstruct section structure.
    """

    output_pages = []

    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):

            has_two_columns, boundary = detect_two_columns(page)

            blocks = page.get_text("blocks")

            if not has_two_columns:
                # Preserve ordinary reading order on pages
                # that do not appear to have two columns.
                page_text = page.get_text("text", sort=True)

                output_pages.append(
                    f"=== PAGE {page_number} ===\n{page_text}"
                )
                continue

            left_blocks = []
            right_blocks = []
            spanning_blocks = []

            for block in blocks:
                x0, y0, x1, y1, text = block[:5]
                text = text.strip()

                if not text:
                    continue

                width = x1 - x0

                # Keep wide blocks and blocks crossing the
                # estimated boundary as separate content.
                if (
                    width >= page.rect.width * 0.75
                    or (x0 < boundary < x1)
                ):
                    spanning_blocks.append((y0, text))

                elif (x0 + x1) / 2 < boundary:
                    left_blocks.append((y0, text))

                else:
                    right_blocks.append((y0, text))

            left_blocks.sort(key=lambda item: item[0])
            right_blocks.sort(key=lambda item: item[0])
            spanning_blocks.sort(key=lambda item: item[0])

            def join_blocks(items):
                return "\n\n".join(
                    text for _, text in items
                )

            page_text = (
                f"=== PAGE {page_number}: TWO-COLUMN LAYOUT ===\n\n"
                "=== FULL-WIDTH OR SPANNING CONTENT ===\n"
                f"{join_blocks(spanning_blocks)}\n\n"
                "=== LEFT COLUMN: TOP TO BOTTOM ===\n"
                f"{join_blocks(left_blocks)}\n\n"
                "=== RIGHT COLUMN: TOP TO BOTTOM ===\n"
                f"{join_blocks(right_blocks)}"
            )

            output_pages.append(page_text)

    return "\n\n".join(output_pages)


# --------------------------------------------------
# Extract PDF content into LLM-friendly Markdown
# --------------------------------------------------

def extract_pdf_text(file_path: str | Path) -> str:
    file_path = Path(file_path)

    # Use coordinate-based extraction if any page appears
    # to have two columns.
    with pymupdf.open(file_path) as pdf:
        has_columns = any(
            detect_two_columns(page)[0]
            for page in pdf
        )

    if has_columns:
        text = extract_pdf_columns(file_path)
    else:
        text = pymupdf4llm.to_markdown(str(file_path))

    if not text.strip():
        raise ValueError(
            f"No text extracted from {file_path.name}"
        )

    return text


# --------------------------------------------------
# Extract text from DOCX
# --------------------------------------------------

def extract_docx_text(file_path: Path) -> str:
    document = Document(file_path)
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


# --------------------------------------------------
# Read TXT files using common encodings
# --------------------------------------------------

def extract_txt_text(file_path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError(
        f"Could not decode text file: {file_path}"
    )


# --------------------------------------------------
# Extract text from any supported resume format
# --------------------------------------------------

def extract_resume_text(file_path: str | Path) -> str:
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        text = extract_pdf_text(file_path)

    elif extension == ".docx":
        text = extract_docx_text(file_path)

    else:
        text = extract_txt_text(file_path)

    text = text.strip()

    if not text:
        raise ValueError(
            f"No text extracted from {file_path.name}. "
            "The file may be scanned or empty."
        )

    return text